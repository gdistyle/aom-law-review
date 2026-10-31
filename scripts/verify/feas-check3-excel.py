import json, re, zipfile, subprocess, sys, os, shutil
from openpyxl import load_workbook
mode=sys.argv[1]
src=f'feas_{mode}.xlsx'; nov=f'v3/nov_{mode}.xlsx'
zin=zipfile.ZipFile(src); zout=zipfile.ZipFile(nov,'w',zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    d=zin.read(it.filename)
    if it.filename.startswith('xl/worksheets/'):
        s=d.decode(); s=re.sub(r'(<f>[^<]*</f>)<v>[^<]*</v>',r'\1',s); d=s.encode()
    zout.writestr(it,d)
zout.close()
os.makedirs('v3/out',exist_ok=True)
subprocess.run(['soffice','--headless','--convert-to','xlsx','--outdir','v3/out',nov],capture_output=True,timeout=120)
wb=load_workbook(f'v3/out/nov_{mode}.xlsx',data_only=True); ws=wb['사업수지']
j=json.load(open(f'feas_{mode}.json')); at=j['at']; res=j['res']
keymap={'land':'land','landTax':'landTax','build':'build','design':'design','buildTax':'buildTax','other':'other','cont':'cont','subtotal':'subtotal','loan':'loan','interest':'interest','revenue':'revenue','saleCost':'saleCost','total':'total','profit':'profit','margin':'margin','equity':'equity','roe':'roe','breakEven':'breakEvenUnit','grossRent':'grossRent','opex':'opex','noi':'noi','deposit':'deposit','netInvest':'netInvest','yield':'yield','payback':'payback','value':'value','devProfit':'devProfit','irr':'irr','npv':'npv','perM2':'perM2'}
bad=0;n=0
for k,rk in keymap.items():
    if k not in at: continue
    xv=ws[at[k]].value; jv=res.get(rk)
    n+=1
    if jv is None and (xv in (None,'')): continue
    if isinstance(xv,str) or xv is None or jv is None: print('MISMATCH',k,xv,jv); bad+=1; continue
    tol=1e-6 if abs(jv)<10 else 0.5
    if abs(xv-jv)>tol: print('MISMATCH',k,xv,jv); bad+=1
v=ws[at['verdict']].value if 'verdict' in at else None
sens=j.get('sens')
if sens is None:
    import subprocess
if 'sens' in at and sens is None:
    pass
if 'sens' in at:
    sv=j.get('sens')
    if sv:
        for i in range(3):
            for k in range(3):
                xv=ws[at['sens'][i][k]].value; jv=sv[i][k]; n+=1
                if jv is None and xv in (None,''): continue
                if not isinstance(xv,(int,float)) or abs(xv-jv)>1e-9: print('SENS',i,k,xv,jv); bad+=1
print(mode,'cells',n,'mismatch',bad,'verdict_excel',v,'verdict_js',res['verdict']['label'])
