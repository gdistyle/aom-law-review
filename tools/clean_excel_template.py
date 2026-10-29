"""설계개요 빈양식 정리: 설계개요·토지내역 2개 시트만, 표 바깥 보조 계산·노란 칠 제거, 층별면적표 수식 정리.
실행: python3 tools/clean_excel_template.py  → public/templates/설계개요_양식.xlsx"""
import openpyxl, copy
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string as ci, get_column_letter as cl
SRC='public/templates/설계개요_빈양식.xlsx'; OUT='public/templates/설계개요_양식.xlsx'
wb=openpyxl.load_workbook(SRC)
for n in list(wb.sheetnames):
    if n not in ('ALT-4(jsl)','토지내역'): del wb[n]
ws=wb['ALT-4(jsl)']; ws.title='설계개요'; land=wb['토지내역']
for s in (ws,land):
    for row in s.iter_rows():
        for c in row:
            if type(c).__name__!='MergedCell': c.comment=None
NOFILL=PatternFill(fill_type=None)
KEEP={'FFC0C0C0','FFFFFFFF'}
def clear(c,fill=True):
    if type(c).__name__=='MergedCell': return
    c.value=None
    if fill: c.fill=NOFILL
# 1) 표 바깥(인쇄 영역 B5:AD38 밖) 전부 비우기
for row in ws.iter_rows(min_row=1,max_row=max(ws.max_row,89),max_col=max(ws.max_column,46)):
    for c in row:
        if c.row<5 or c.row>38 or c.column<2 or c.column>30: clear(c)
for r in range(39,90): ws.row_dimensions[r].hidden=False
for mr in list(ws.merged_cells.ranges):
    if mr.min_row<5 or mr.min_row>38 or mr.min_col>30: ws.unmerge_cells(str(mr))
# 2) 노랑·주황·연두 칠 제거(머리글 회색만 유지)
for row in ws.iter_rows(min_row=5,max_row=38,min_col=2,max_col=30):
    for c in row:
        f=c.fill
        if f and f.fill_type=='solid' and (f.fgColor.type!='rgb' or f.fgColor.rgb not in KEEP): c.fill=NOFILL
# 3) 원래 예시값·보조 수식 제거
for a in ['C6','D9','E9','F9','D10','E10','F10','D11','E11','F11','D12','E12','F12','E13','F13','E14','F14','D16',
          'D26','E26','G26','D27','E27','G27','G30','P22','P23','P24','AB25','AB27','AB28','AB30']:
    clear(ws[a],fill=False)
ws['D26'].value=None
# 층별면적표 각 행(지상 9~28, 지하 30~35)은 사이트가 사용하는 층만 채움 → 행 수식·라벨 비움
for r in list(range(9,29))+list(range(30,36)):
    for col in range(ci('I'),ci('AA')+1): clear(ws.cell(r,col),fill=False)
# 4) 합계 수식(9~28행 전체 합산으로 정정)
F={'J29':'=SUM(J9:J28)','K29':'=J29*0.3025','L29':'=SUM(L9:L28)','M29':'=L29*0.3025','N29':'=SUM(N9:N28)','O29':'=N29*0.3025',
   'P29':'=SUM(P9:P28)','Q29':'=P29*0.3025','AA29':'=P29',
   'J36':'=SUM(J30:J35)','K36':'=J36*0.3025','L36':'=SUM(L30:L35)','M36':'=L36*0.3025','N36':'=SUM(N30:N35)','O36':'=N36*0.3025',
   'P36':'=SUM(P30:P35)','Q36':'=P36*0.3025','AA36':'=P36',
   'J37':'=J29+J36','K37':'=J37*0.3025','L37':'=L29+L36','M37':'=L37*0.3025','N37':'=N29+N36','O37':'=N37*0.3025','P37':'=P29+P36','Q37':'=P37*0.3025','AA37':'=P37',
   'E15':'=SUM(E9:E12)-E13-E14','F15':'=E15*0.3025','E21':'=P29','F21':'=E21*0.3025','E22':'=P36','F22':'=E22*0.3025','E23':'=E21+E22','F23':'=E23*0.3025',
   'E20':'=D20*0.3025','D24':'=IF(D20="","",D20/$E$15)','D25':'=IF(P29=0,"",P29/$E$15)',
   'D34':'=E15','F34':'=IF(E34="","",D34*E34)','D36':'=D34','E36':'=E34','F36':'=IF(E36="","",D36*E36)'}
for k,v in F.items(): ws[k].value=v
for k in ['S29','T29','U29','V29','W29','X29','Y29','Z29','R36','S36','T36','U36','V36','W36','X36','Y36','Z36','R37','S37','T37','U37','V37','W37','X37','Y37','Z37','R38','S38','T38','U38','V38','W38','X38','Y38','Z38','AA38']:
    clear(ws[k],fill=False)
# 서식 통일
ACC='_-* #,##0.00_-;\\-* #,##0.00_-;_-* \\-_-;_-@_-'
for r in list(range(9,29))+list(range(30,36)):
    for col in 'JKLMNOPQ': ws[f'{col}{r}'].number_format=ACC
    ws[f'I{r}'].number_format='@'
for a in ['J29','K29','L29','M29','N29','O29','P29','Q29','J36','K36','L36','M36','N36','O36','P36','Q36','J37','K37','L37','M37','N37','O37','P37','Q37']: ws[a].number_format='#,##0.00_ '
ws['E26'].number_format='#,##0.00"㎡"'; ws['G26'].number_format='0" 대"'; ws['G30'].number_format='0" 대"'; ws['G31'].number_format='0" 대"'; ws['G32'].number_format='0" 대"'
ws['D9'].number_format='@'
for a in ['D9','D10','D11','D12']: ws[a].number_format='@'
ws['C6'].number_format='@'; ws['AD5'].number_format='@'
# 법정 최대(적용 면적) 요약표 — 표 위 1~4행
from openpyxl.styles import Border, Side, Font, Alignment as AL
th=Side(style='thin'); BD=Border(left=th,right=th,top=th,bottom=th); GRAY=PatternFill('solid',fgColor='FFC0C0C0')
for r in range(1,5):
    for col in 'BCDEFGH':
        c=ws[f'{col}{r}']; c.border=BD; c.font=Font(name='맑은 고딕',size=11,bold=(r==1)); c.alignment=AL(horizontal='center',vertical='center')
        c.fill=GRAY if r==1 else NOFILL
    ws.merge_cells(f'B{r}:C{r}'); ws.merge_cells(f'F{r}:H{r}')
    ws[f'F{r}'].alignment=AL(horizontal='center' if r==1 else 'left',vertical='center')
    ws.row_dimensions[r].height=20
for k,v in {'B1':'법정 최대 (적용 기준)','D1':'㎡','E1':'평','F1':'산식 · 근거','B2':'건축면적','B3':'연면적','B4':'지상 층수'}.items(): ws[k].value=v
for r in (2,3): ws[f'D{r}'].number_format='#,##0.00'; ws[f'E{r}'].number_format='#,##0.00'; ws[f'D{r}'].alignment=AL(horizontal='right',vertical='center'); ws[f'E{r}'].alignment=AL(horizontal='right',vertical='center')
ws['D4'].number_format='0"층"'; ws.merge_cells('D4:E4')
ws.print_area='B1:AD38'
for a in ('D27','D28','D35'):
    c=ws[a]; c.alignment=AL(horizontal='left',vertical='center'); f=copy.copy(c.font); c.font=Font(name=f.name,size=9,color='FF595959')
for a in ('D24','D25','F24','F25','E34','E36'):
    f=ws[a].font; ws[a].font=Font(name=f.name,size=f.sz,bold=f.b)
# 토지내역: 표(A1:I29) 밖 보조 계산 제거, 평 환산·평당가 수식 정리
for row in land.iter_rows(min_row=30,max_row=max(land.max_row,40)):
    for c in row: clear(c)
for r in range(3,29):
    for col in 'ABCDEFGHI': land[f'{col}{r}'].value=None
land['C29']='=SUM(C3:C28)'; land['D29']='=SUM(D3:D28)'
land.print_area='A1:I29'
for row in land.iter_rows(min_row=1,max_row=29,max_col=9):
    for c in row:
        f=c.fill
        if f and f.fill_type=='solid' and (f.fgColor.type!='rgb' or f.fgColor.rgb not in KEEP): c.fill=NOFILL
for r in range(3,29):
    land[f'F{r}'].number_format='#,##0'; land[f'G{r}'].number_format='#,##0'; land[f'H{r}'].number_format='#,##0'
    land[f'I{r}'].alignment=copy.copy(land[f'A{r}'].alignment)
land.page_setup.orientation='portrait'; land.sheet_properties.pageSetUpPr.fitToPage=True; land.page_setup.fitToWidth=1; land.page_setup.fitToHeight=1
from openpyxl.styles import Alignment
a=copy.copy(ws['D16'].alignment); ws['D16'].alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
ws.row_dimensions[16].height=34
wb.active=0
wb.save(OUT)
print('ok', wb.sheetnames)
